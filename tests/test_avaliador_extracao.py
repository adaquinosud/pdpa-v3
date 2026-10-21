"""Extração do Avaliador com documentos SINTÉTICOS gerados no próprio teste (nenhum
documento real entra na suíte)."""

from __future__ import annotations

import io
import zipfile

import pytest

from src.avaliador import extracao
from src.avaliador.extracao import (
    Documento,
    DocumentoRecusado,
    Pagina,
    conferir_trecho,
    detectar_formato,
    extrair,
)

TEXTO_LONGO = "Entregamos o painel mensal até o quinto dia útil de cada mês."


def _pptx() -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    branco = prs.slide_layouts[6]
    # 1: caixa de texto simples + nota do apresentador
    s1 = prs.slides.add_slide(branco)
    s1.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text_frame.text = TEXTO_LONGO
    s1.notes_slide.notes_text_frame.text = "Lembrar de citar o prazo de resposta de quatro horas."
    # 2: texto só dentro de um GRUPO aninhado
    s2 = prs.slides.add_slide(branco)
    grupo = s2.shapes.add_group_shape()
    sub = grupo.shapes.add_group_shape()
    sub.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text_frame.text = (
        "Texto escondido dentro de um grupo aninhado de formas."
    )
    # 3: texto só numa TABELA
    s3 = prs.slides.add_slide(branco)
    tab = s3.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(6), Inches(2)).table
    tab.cell(0, 0).text = "Nível de serviço"
    tab.cell(0, 1).text = "Resposta em até quatro horas úteis"
    # 4: vazio (só imagem, na vida real)
    prs.slides.add_slide(branco)
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _pdf(paginas) -> bytes:
    """PDF mínimo escrito à mão (Helvetica/WinAnsi), uma linha de texto por página.
    Sem WeasyPrint: as libs nativas dele não existem fora da imagem Docker."""
    n = len(paginas)
    objs = {1: b"<< /Type /Catalog /Pages 2 0 R >>"}
    kids = " ".join(f"{3 + 2 * i} 0 R" for i in range(n))
    objs[2] = f"<< /Type /Pages /Kids [{kids}] /Count {n} >>".encode()
    fonte = 3 + 2 * n
    for i, t in enumerate(paginas):
        pag, cont = 3 + 2 * i, 4 + 2 * i
        objs[pag] = (
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {cont} 0 R "
            f"/Resources << /Font << /F1 {fonte} 0 R >> >> >>"
        ).encode()
        esc = t.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        fluxo = b"BT /F1 10 Tf 50 700 Td (" + esc.encode("cp1252") + b") Tj ET" if t else b""
        objs[cont] = b"<< /Length %d >>\nstream\n" % len(fluxo) + fluxo + b"\nendstream"
    objs[fonte] = (
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>"
    )
    out = bytearray(b"%PDF-1.4\n")
    offs = {}
    for k in sorted(objs):
        offs[k] = len(out)
        out += b"%d 0 obj\n" % k + objs[k] + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1)
    for k in sorted(objs):
        out += b"%010d 00000 n \n" % offs[k]
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objs) + 1, xref)
    return bytes(out)


def test_pptx_le_grupo_tabela_e_nota_e_conta_pagina_vazia():
    doc = extrair(_pptx())
    assert doc.formato == "pptx" and doc.n_paginas == 4
    assert TEXTO_LONGO in doc.paginas[0].texto
    assert "quatro horas" in doc.paginas[0].notas  # nota fica separada do corpo
    assert "grupo aninhado" in doc.paginas[1].texto
    assert "Resposta em até quatro horas úteis" in doc.paginas[2].texto
    assert doc.paginas_sem_texto == [4]


def test_nota_vai_marcada_no_prompt():
    prompt = extrair(_pptx()).para_prompt()
    assert "=== PÁGINA 1 ===" in prompt
    assert "nota do apresentador" in prompt
    assert "[página sem texto legível]" in prompt


def test_pdf_por_pagina():
    doc = extrair(_pdf([TEXTO_LONGO, "Segunda página com conteúdo suficiente aqui.", ""]))
    assert doc.formato == "pdf"
    assert doc.n_paginas == 3
    assert "quinto dia útil" in doc.paginas[0].texto
    assert doc.paginas_sem_texto == [3]


def test_json_ida_e_volta():
    doc = extrair(_pptx())
    volta = Documento.de_json(doc.para_json())
    assert [(p.numero, p.texto, p.notas) for p in volta.paginas] == [
        (p.numero, p.texto, p.notas) for p in doc.paginas
    ]


def test_formato_pelo_conteudo_e_nao_pela_extensao():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("word/document.xml", "<w/>")
    assert detectar_formato(buf.getvalue()) == "docx"  # pelo conteúdo, não por ".docx"
    with pytest.raises(DocumentoRecusado, match="não reconhecido"):
        detectar_formato(b"\xd0\xcf\x11\xe0 ppt antigo")


def test_pdf_com_senha_recusado():
    from pypdf import PdfReader, PdfWriter

    w = PdfWriter()
    for p in PdfReader(io.BytesIO(_pdf([TEXTO_LONGO]))).pages:
        w.add_page(p)
    w.encrypt("segredo")
    buf = io.BytesIO()
    w.write(buf)
    with pytest.raises(DocumentoRecusado, match="senha"):
        extrair(buf.getvalue())


def test_so_imagem_recusado():
    with pytest.raises(DocumentoRecusado, match="só imagem"):
        extrair(_pdf(["", ""]))


# A forma real (resumo executivo do próprio app, WeasyPrint, pypdf 6.20.1, 14/10): um
# nulo antes de cada letra. O arquivo real não entra na suíte; a forma, sim.
LIXO = "".join("\x00" + c for c in "DIAGNÓSTICO PONTUAL · PDPA Grupo Exemplo " * 3)


def _com_paginas(monkeypatch, textos):
    doc = Documento("pdf", [Pagina(i, t) for i, t in enumerate(textos, start=1)])
    monkeypatch.setattr(extracao, "_extrair_pdf", lambda dados: doc)
    return _pdf(["x"])


def test_pdf_com_texto_ilegivel_e_recusado_e_nao_vira_leitura(monkeypatch):
    dados = _com_paginas(monkeypatch, [LIXO, LIXO])
    with pytest.raises(DocumentoRecusado, match="ilegível: em 2 de 2 páginas.*cole o texto"):
        extrair(dados)


def test_pagina_ilegivel_no_meio_nao_vai_ao_modelo_e_e_listada(monkeypatch):
    doc = extrair(_com_paginas(monkeypatch, [TEXTO_LONGO, LIXO, TEXTO_LONGO]))
    assert doc.paginas_sem_texto == [2]
    assert "\x00" not in doc.para_prompt()
    assert "=== PÁGINA 2 ===\n[página sem texto legível]" in doc.para_prompt()


def test_nulo_solto_em_pagina_curta_nao_e_ilegivel(monkeypatch):
    # 1 nulo numa página de 10 letras passa de 2%, mas não do mínimo absoluto (20)
    texto = TEXTO_LONGO + "\n\x00"
    doc = extrair(_com_paginas(monkeypatch, [texto, "try/catch\n\x00"]))
    assert doc.paginas[0].texto == texto


def test_limites(monkeypatch):
    dados = _pptx()
    monkeypatch.setattr(extracao, "LIMITE_PAGINAS", 3)
    with pytest.raises(DocumentoRecusado, match="limite é 3"):
        extrair(dados)
    monkeypatch.setattr(extracao, "LIMITE_PAGINAS", 120)
    monkeypatch.setattr(extracao, "LIMITE_CARACTERES", 50)
    with pytest.raises(DocumentoRecusado, match="texto demais"):
        extrair(dados)
    monkeypatch.setattr(extracao, "LIMITE_BYTES", 10)
    with pytest.raises(DocumentoRecusado, match="maior que"):
        extrair(dados)


def _doc():
    return Documento(
        "pdf",
        [
            Pagina(1, "Atendimento:   de segunda a sexta,\n das 8h às 18h."),
            Pagina(2, "O gerente responde em “até 4 horas” — sempre."),
        ],
    )


def test_citacao_exata_e_com_espaco_diferente():
    assert conferir_trecho(_doc(), 1, "de segunda a sexta, das 8h às 18h") == (True, 1)


def test_citacao_na_pagina_errada_e_achada_e_corrigida():
    assert conferir_trecho(_doc(), 1, 'responde em "até 4 horas" - sempre') == (True, 2)


def test_citacao_inexistente_ou_parafraseada_nao_confere():
    assert conferir_trecho(_doc(), 1, "atendimento em horário comercial") == (False, None)
    assert conferir_trecho(_doc(), 1, "DE SEGUNDA A SEXTA, DAS 8H") == (False, None)


def test_citacao_curta_demais_nao_prova_nada():
    assert conferir_trecho(_doc(), 1, "sexta") == (False, None)


def test_texto_colado_vira_documento_de_uma_pagina():
    from src.avaliador.extracao import documento_de_texto

    doc = documento_de_texto("Cliente: com quem eu falo?\nVendedor: comigo, das 8h às 18h.")
    assert doc.formato == "texto" and doc.n_paginas == 1
    assert "=== TEXTO COLADO ===" in doc.para_prompt()
    assert conferir_trecho(doc, 1, "comigo, das 8h às 18h") == (True, 1)
    with pytest.raises(DocumentoRecusado, match="vazio"):
        documento_de_texto("   ")


# ── costuras do modelo (21/10): emenda e elisão ───────────────────────────

CARTAO = "Consultoria de CX\nDiagnóstico tradicional entrega um relatório em seis semanas."
PARECER = (
    "O parecer é fechado em si. Quem quiser seguir pode contratar a Leitura Contínua, "
    "que acompanha mês a mês. As duas são contratadas à parte."
)


def _costuras():
    return Documento("pdf", [Pagina(1, CARTAO), Pagina(2, PARECER), Pagina(3, "Outra coisa.")])


def test_citacao_literal_continua_como_o_modelo_escreveu():
    assert extracao.localizar_trecho(_costuras(), 2, "fechado em si. Quem quiser") == (
        2,
        "fechado em si. Quem quiser",
        "",
    )


def test_emenda_pontuacao_no_lugar_da_quebra_de_linha():
    for junta in (": ", " - ", ". ", " | "):
        r = extracao.localizar_trecho(
            _costuras(), 1, f"Consultoria de CX{junta}Diagnóstico tradicional"
        )
        assert r == (1, "Consultoria de CX Diagnóstico tradicional", "emenda"), junta


def test_emenda_nao_aceita_palavra_nova_nem_pontuacao_onde_a_peca_nao_quebra():
    doc = _costuras()
    assert conferir_trecho(doc, 1, "Consultoria de RH: Diagnóstico tradicional") == (False, None)
    assert conferir_trecho(doc, 1, "Consultoria de CX, Diagnóstico tradicional") == (False, None)
    assert conferir_trecho(doc, 1, "Diagnóstico tradicional: entrega um relatório") == (False, None)


def test_elisao_dois_pedacos_literais_da_mesma_pagina_na_ordem():
    citado = (
        "O parecer é fechado em si. Quem quiser seguir pode contratar a Leitura Contínua"
        "... As duas são contratadas à parte"
    )
    achada, literal, ajuste = extracao.localizar_trecho(_costuras(), 2, citado)
    assert (achada, ajuste) == (2, "elisao")
    assert literal == (
        "O parecer é fechado em si. Quem quiser seguir pode contratar a Leitura Contínua"
        " […] As duas são contratadas à parte"
    )
    for marca in ("…", "[...]", "(…)"):
        assert conferir_trecho(
            _costuras(), 2, f"O parecer é fechado em si {marca} As duas são contratadas"
        )[0]


def test_elisao_recusada_fora_de_ordem_entre_paginas_curta_ou_inventada():
    doc = _costuras()
    assert not conferir_trecho(
        doc, 2, "As duas são contratadas à parte... O parecer é fechado em si"
    )[0]
    assert not conferir_trecho(
        doc, 2, "O parecer é fechado em si... Diagnóstico tradicional entrega"
    )[0]
    assert not conferir_trecho(doc, 2, "O parecer é fechado em si... à parte")[0]
    assert not conferir_trecho(doc, 2, "O parecer é fechado em si... e custa pouco mais de nada")[0]


def test_reticencias_na_ponta_do_trecho():
    r = extracao.localizar_trecho(_costuras(), 2, "...pode contratar a Leitura Contínua")
    assert r == (2, "pode contratar a Leitura Contínua", "elisao")


def test_motor_grava_o_trecho_da_peca_e_guarda_o_que_o_modelo_escreveu():
    from src.avaliador.motor import _evidencias

    citado = "Consultoria de CX: Diagnóstico tradicional entrega"
    [ev] = _evidencias(_costuras(), [{"pagina": 1, "trecho": citado}])
    assert ev["conferido"] and ev["trecho"] == "Consultoria de CX Diagnóstico tradicional entrega"
    assert ev["ajuste"] == "emenda" and ev["trecho_citado"] == citado
    [ev] = _evidencias(_costuras(), [{"pagina": 1, "trecho": "Diagnóstico tradicional entrega"}])
    assert "ajuste" not in ev and "trecho_citado" not in ev
