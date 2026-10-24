"""Serviço do Avaliador com cliente LLM DUBLADO (nenhuma chamada paga) e banco de teste."""

from __future__ import annotations

import io
import json
import logging
import re
from datetime import datetime, timedelta
from types import SimpleNamespace

import pytest

from src.avaliador import servico
from src.avaliador.servico import AvaliacaoRecusada
from src.models.avaliador import STATUS_FALHOU, STATUS_LENDO, STATUS_PRONTA, AvaliadorAnalise
from src.models.empresa import Empresa

SENTINELA = "SENTINELA-DOCUMENTO-CONFIDENCIAL-9F3A"
DESCRICAO = "Proposta comercial para renovar o contrato de um cliente de TI"


def config_avaliador(monkeypatch, **valores):
    """Troca valores de config SÓ para o serviço do Avaliador (o único que os lê).
    Patch na instância devolvida, não na classe Config — imune a outros testes."""
    from src.config import get_config as real

    def fake():
        c = real()
        for k, v in valores.items():
            setattr(c, k, v)
        return c

    monkeypatch.setattr(servico, "get_config", fake)


def _pptx() -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    for texto in (
        f"Proposta comercial para a Empresa X. {SENTINELA}",
        "Atendimento de segunda a sexta das 8h às 18h pelo gerente de conta.",
        "Somos referência no setor há vinte anos, com clientes em todo o país.",
    ):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        s.shapes.add_textbox(Inches(1), Inches(1), Inches(7), Inches(2)).text_frame.text = texto
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def veredito_de(cod, ver="nao", trecho=None, pagina=1, onde=None, deveria="Dizer como e quando."):
    return {
        "subpilar": cod,
        "veredito": ver,
        "onde": onde or ("Esta conversa não existe na peça." if ver == "nao" else "Está na peça."),
        "deveria_ter": "" if ver == "sim" and deveria == "Dizer como e quando." else deveria,
        "evidencias": [{"pagina": pagina, "trecho": trecho}] if trecho else [],
    }


def resposta_v20(vereditos=None, situar=None, lastro="Padrão de teste. Comece pela Precisão."):
    """Leitura no formato 2, dublada: D1 "sim" e A1 "parcial" com trecho da peça do
    _pptx(), os outros dez "nao". ``vereditos`` troca subpilares; ``lastro=None`` faz a
    etapa final responder sem JSON (falha)."""
    from src.api.painel import SUBPILARES_ORDEM

    ver = {c: veredito_de(c) for c in SUBPILARES_ORDEM}
    ver["D1"] = veredito_de(
        "D1", "sim", "segunda a sexta das 8h às 18h", 2, onde="Horário declarado."
    )
    ver["A1"] = veredito_de(
        "A1",
        "parcial",
        "Somos referência no setor há vinte anos",
        3,
        onde="Referência afirmada sem quem a cite.",
        deveria="Quem cita a empresa como referência.",
    )
    ver.update(vereditos or {})
    sit = {
        "tipo_peca": "Proposta comercial.",
        "leitor": "Quem decide sozinho.",
        "situacao": "Lida antes da reunião.",
        "consequencia_cobertura": "",
        "achados": [],
        "contradicoes": [],
    }
    sit.update(situar or {})
    return {"v20": True, "situar": sit, "vereditos": ver, "lastro": lastro}


HORARIO = (2, "segunda a sexta das 8h às 18h")
REFERENCIA = (3, "Somos referência no setor há vinte anos")


def resposta_v25(presentes=None, achados=None, contradicoes=None, achados_falham=False):
    """Leitura no formato 3, dublada: ``presentes`` é código da marca → (página, trecho);
    o resto não está. Padrão: D1.01 e D1.04 com o MESMO trecho do horário (p. 2) e A1.02
    com a frase de referência (p. 3) do _pptx(). ``achados_falham`` faz a etapa dos
    achados responder sem JSON (a leitura falha)."""
    if presentes is None:
        presentes = {"D1.01": HORARIO, "D1.04": HORARIO, "A1.02": REFERENCIA}
    return {
        "v25": True,
        "presentes": presentes,
        "achados": achados or [],
        "contradicoes": contradicoes or [],
        "achados_falham": achados_falham,
    }


def _etapa_de_marcas(pedido, presentes):
    codigos = re.search(r"destas marcas: ([A-Za-z0-9., ]+)\. Nenhuma", pedido).group(1)
    out = []
    for c in codigos.split(", "):
        if c in presentes:
            trechos = presentes[c] if isinstance(presentes[c], list) else [presentes[c]]
            ev = [{"pagina": pag, "trecho": tr} for pag, tr in trechos]
            out.append({"codigo": c, "evidencias": ev, "esta": True})
        else:
            out.append({"codigo": c, "evidencias": [], "esta": False})
    return {"marcas": out}


def _verificacao(entrada: str, exclui) -> dict:
    """A resposta da verificação dos trechos (24/10): "passa" para todo trecho, menos os que
    contêm um dos textos de ``exclui`` ({texto: tipo}), que saem com o tipo dado."""
    out = []
    for n, trecho in re.findall(r"^\[(\d+)\]\ntrecho: “(.*)”$", entrada, re.M):
        tipo = next((t for k, t in (exclui or {}).items() if k in trecho), None)
        out.append({"n": int(n), "decisao": "exclui" if tipo else "passa", "tipo": tipo,
                    "por_que": "dublê"})  # fmt: skip
    return {"trechos": out}


def _leitura(chamadas):
    """As chamadas da leitura, sem a da verificação dos trechos que vem no fim (24/10)."""
    return [c for c in chamadas if not c["messages"][0]["content"][-1]["text"].startswith(
        "Decida cada trecho")]  # fmt: skip


class Dublê:
    def __init__(self, respostas=None, quebra=False, exclui=None):
        self.respostas = list(respostas or [resposta_v25()])
        self.quebra = quebra
        self.exclui = exclui  # a verificação dos trechos: {texto: tipo} que sai do Word
        self.chamadas = []
        self.messages = self

    def create(self, **kw):
        self.chamadas.append(kw)
        if self.quebra:
            raise ConnectionError("rede caiu")
        conteudo = kw["messages"][0]["content"]
        if isinstance(conteudo, list) and conteudo[-1]["text"].startswith("Decida cada trecho"):
            texto = json.dumps(_verificacao(conteudo[0]["text"], self.exclui), ensure_ascii=False)
            return SimpleNamespace(
                content=[SimpleNamespace(type="text", text=texto)],
                usage=SimpleNamespace(input_tokens=1000, output_tokens=200,
                                      cache_creation_input_tokens=0, cache_read_input_tokens=0),
            )  # fmt: skip
        r = self.respostas[0]
        if isinstance(conteudo, list) and isinstance(r, dict) and r.get("v25"):
            # formato 3: responde pela etapa pedida no fim da mensagem; a resposta inteira
            # é consumida na etapa dos achados (a última)
            pedido = conteudo[-1]["text"]
            if pedido.startswith("ETAPA: A PEÇA INTEIRA"):
                if r["achados_falham"]:
                    r = "sem json"
                else:
                    r = {k: r[k] for k in ("achados", "contradicoes")}
                    self.respostas.pop(0)
            else:
                r = _etapa_de_marcas(pedido, r["presentes"])
        elif isinstance(conteudo, list) and isinstance(r, dict) and r.get("v20"):
            # formato 2: responde pela etapa pedida no fim da mensagem; a resposta inteira
            # é consumida na etapa final válida (a inválida se repete na nova tentativa)
            pedido = conteudo[-1]["text"]
            if pedido.startswith("ETAPA 1"):
                r = r["situar"]
            elif pedido.startswith("ETAPA: A PEÇA INTEIRA"):
                r = {k: r["situar"][k] for k in ("achados", "contradicoes")}
            elif pedido.startswith("ETAPA FINAL"):
                if r["lastro"] is None:
                    r = "sem json"  # etapa final inválida: a leitura fica sem o texto
                else:
                    r = {"lastro": r["lastro"]}
                    self.respostas.pop(0)
            else:
                cods = re.search(r"subpilares ([A-Za-z0-9, ]+)\.", pedido).group(1).split(", ")
                r = {"vereditos": [r["vereditos"][c] for c in cods]}
        else:
            self.respostas.pop(0)
        texto = r if isinstance(r, str) else json.dumps(r, ensure_ascii=False)
        return SimpleNamespace(
            content=[SimpleNamespace(type="text", text=texto)],
            usage=SimpleNamespace(
                input_tokens=3000,
                output_tokens=800,
                cache_creation_input_tokens=0,
                cache_read_input_tokens=0,
            ),
        )


@pytest.fixture
def empresa(db_session):
    e = Empresa(nome="Empresa Teste")
    db_session.add(e)
    db_session.commit()
    return e


@pytest.fixture
def teto(monkeypatch):
    config_avaliador(monkeypatch, AVALIADOR_TETO_MENSAL_USD=10.0)


def _criar(db_session, empresa, **kw):
    if not kw:
        kw = {"nome": "proposta.pptx", "dados": _pptx()}
    kw.setdefault("cliente_nomeado", False)
    a = servico.criar_analise(db_session, empresa.id, None, DESCRICAO, **kw)
    db_session.commit()
    return a.id


def _rodar(db_session, aid, cli):
    servico.processar(aid, client=cli)
    db_session.expire_all()
    return db_session.get(AvaliadorAnalise, aid)


def test_teto_zero_desliga_antes_de_gravar(db_session, empresa):
    with pytest.raises(AvaliacaoRecusada, match="desligado"):
        servico.criar_analise(
            db_session, empresa.id, None, DESCRICAO, cliente_nomeado=False, dados=_pptx()
        )
    assert db_session.query(AvaliadorAnalise).count() == 0


def test_teto_estourado_recusa(db_session, empresa, monkeypatch):
    config_avaliador(monkeypatch, AVALIADOR_TETO_MENSAL_USD=0.0001)
    with pytest.raises(AvaliacaoRecusada, match="passaria do teto"):
        servico.criar_analise(
            db_session, empresa.id, None, DESCRICAO, cliente_nomeado=False, dados=_pptx()
        )


def test_descricao_e_uma_entrada_sao_obrigatorias(db_session, empresa, teto):
    with pytest.raises(AvaliacaoRecusada, match="o que é a peça"):
        servico.criar_analise(
            db_session, empresa.id, None, "  ", cliente_nomeado=False, dados=_pptx()
        )
    with pytest.raises(AvaliacaoRecusada, match="um só"):
        servico.criar_analise(db_session, empresa.id, None, DESCRICAO, cliente_nomeado=False)
    with pytest.raises(AvaliacaoRecusada, match="um só"):
        servico.criar_analise(
            db_session,
            empresa.id,
            None,
            DESCRICAO,
            cliente_nomeado=False,
            dados=_pptx(),
            texto="x" * 50,
        )


def test_caminho_feliz_arquivo(db_session, empresa, teto):
    aid = _criar(db_session, empresa)
    cli = Dublê()
    a = _rodar(db_session, aid, cli)
    assert a.status == STATUS_PRONTA and a.resumo is None and a.erro_motivo is None
    assert a.prompt_versao == "avaliador-v27-marcas"
    # uma por pilar (a última com o bloco C) e os achados; depois, a verificação dos trechos
    assert len(_leitura(cli.chamadas)) == 5 and len(cli.chamadas) == 6
    assert DESCRICAO in cli.chamadas[0]["messages"][0]["content"][0]["text"]
    assert a.custo_usd > 0 and a.origem == "arquivo" and a.formato == "pptx"
    assert servico.formato_da_leitura(a) == servico.FORMATO_MARCAS == servico.FORMATO_ATUAL
    org = servico.organizar_marcas(servico.carregar_regua(), a)
    d1 = org["pilares"][1]["subpilares"][0]  # Disponibilidade → D1
    assert d1["subpilar"].codigo == "D1"
    estados = {x["codigo"]: x["estado"] for x in d1["marcas"]}
    assert estados["D1.01"] == estados["D1.04"] == "esta" and estados["D1.02"] == "nao_esta"
    ev = d1["marcas"][0]["evidencias"][0]
    assert ev["conferido"] and ev["pagina_encontrada"] == 2
    assert org["cobertura"][0] == "Lidas 3 páginas."


def test_texto_colado_le_as_mesmas_marcas(db_session, empresa, teto):
    """19/10: as marcas são as mesmas para toda origem; as bibliotecas de pistas por
    origem eram do veredito e não entram no pedido."""
    conversa = "Cliente: com quem eu falo se der problema?\nVendedor: fica tranquilo, é rápido."
    aid = _criar(db_session, empresa, texto=conversa)
    cli = Dublê()
    a = _rodar(db_session, aid, cli)
    assert a.status == STATUS_PRONTA and a.origem == "texto" and a.formato == "texto"
    assert "Ele perguntou com quem fala" not in cli.chamadas[0]["system"]
    assert "D1.01 Diz por onde se fala com a empresa?" in str(cli.chamadas[1]["messages"])
    assert "=== TEXTO COLADO ===" in cli.chamadas[0]["messages"][0]["content"][0]["text"]


def test_falha_fica_visivel_sem_vazar_texto_e_tentar_de_novo_funciona(
    db_session, empresa, teto, caplog
):
    aid = _criar(db_session, empresa)
    with caplog.at_level(logging.WARNING):
        a = _rodar(db_session, aid, Dublê(quebra=True))
    assert a.status == STATUS_FALHOU and a.erro_motivo == "leitura: ConnectionError"
    assert SENTINELA not in caplog.text
    servico.refazer(db_session, a)
    db_session.commit()
    a = _rodar(db_session, aid, Dublê())
    assert a.status == STATUS_PRONTA


def test_resposta_invalida_duas_vezes_cobra_e_falha(db_session, empresa, teto):
    aid = _criar(db_session, empresa)
    a = _rodar(db_session, aid, Dublê(["lixo", "lixo"]))
    assert a.status == STATUS_FALHOU and "RespostaInvalida" in a.erro_motivo
    assert a.tokens_in == 6000 and a.custo_usd > 0


def test_refazer_com_nova_descricao(db_session, empresa, teto):
    aid = _criar(db_session, empresa)
    a = _rodar(db_session, aid, Dublê())
    servico.refazer(db_session, a, "Apresentação institucional para um evento do setor")
    db_session.commit()
    cli = Dublê()
    a = _rodar(db_session, aid, cli)
    assert a.descricao.startswith("Apresentação institucional")
    assert "Apresentação institucional" in cli.chamadas[0]["messages"][0]["content"][0]["text"]
    assert a.tokens_in == 30000  # as duas leituras somam (5 chamadas de 3000 cada)


def test_concluir_apaga_o_texto_e_depois_nao_se_rele(db_session, empresa, teto):
    aid = _criar(db_session, empresa)
    a = _rodar(db_session, aid, Dublê())
    servico.concluir(a)
    db_session.commit()
    assert a.texto_json is None and a.texto_apagado_motivo == "concluida"
    assert len(json.loads(a.itens_json)["marcas"]) == len(servico.carregar_regua().marcas)
    with pytest.raises(AvaliacaoRecusada, match="já foi apagado"):
        servico.refazer(db_session, a)


def test_housekeeping_marca_orfa_e_apaga_texto_sem_uso(db_session, empresa, teto):
    agora = datetime.utcnow()
    orfa = db_session.get(AvaliadorAnalise, _criar(db_session, empresa))
    orfa.status, orfa.atualizado_em = STATUS_LENDO, agora - timedelta(minutes=30)
    velha = db_session.get(AvaliadorAnalise, _criar(db_session, empresa))
    velha.status, velha.usado_em = STATUS_PRONTA, agora - timedelta(days=8)
    nova = db_session.get(AvaliadorAnalise, _criar(db_session, empresa))
    nova.status, nova.usado_em = STATUS_PRONTA, agora - timedelta(days=6)
    db_session.commit()
    assert servico.housekeeping(db_session, agora) == {"orfas": 1, "textos_apagados": 1}
    assert orfa.status == STATUS_FALHOU and "reinício" in orfa.erro_motivo
    assert velha.texto_json is None and velha.texto_apagado_motivo == "expirada"
    assert nova.texto_json is not None
    assert servico.housekeeping(db_session, agora) == {"orfas": 0, "textos_apagados": 0}


def test_retencao_segue_a_constante_de_config(db_session, empresa, monkeypatch):
    config_avaliador(monkeypatch, AVALIADOR_TETO_MENSAL_USD=10.0, AVALIADOR_DIAS_RETENCAO_TEXTO=2)
    a = db_session.get(AvaliadorAnalise, _criar(db_session, empresa))
    a.status, a.usado_em = STATUS_PRONTA, datetime.utcnow() - timedelta(days=3)
    db_session.commit()
    assert servico.housekeeping(db_session)["textos_apagados"] == 1


def test_comando_do_watchdog_roda(client, empresa):
    r = client.application.test_cli_runner().invoke(args=["avaliador-housekeeping"])
    assert r.exit_code == 0, r.output
    assert "0 interrompida(s)" in r.output and "7 dias sem uso" in r.output


def test_leitura_de_versao_antiga_na_fila_sai_nas_marcas_de_hoje(db_session, empresa, teto):
    """Os formatos 1 e 2 não são mais gerados. Leitura que estivesse na fila numa versão
    deles é lida nas marcas de hoje e passa a levar essa versão."""
    aid = _criar(db_session, empresa)
    db_session.get(AvaliadorAnalise, aid).prompt_versao = "avaliador-v14-pilar"
    db_session.commit()
    cli = Dublê()
    a = _rodar(db_session, aid, cli)
    assert a.status == STATUS_PRONTA and a.prompt_versao == "avaliador-v27-marcas"
    assert servico.formato_da_leitura(a) == servico.FORMATO_MARCAS
    assert '"forte"' not in cli.chamadas[0]["system"] and len(_leitura(cli.chamadas)) == 5


def test_falha_nos_achados_falha_a_leitura_e_cobra(db_session, empresa, teto):
    """As marcas saíram, mas a etapa dos achados falhou duas vezes: a leitura fica FALHOU
    (como no formato 2), com o gasto das seis chamadas."""
    aid = _criar(db_session, empresa)
    a = _rodar(db_session, aid, Dublê([resposta_v25(achados_falham=True)]))
    assert a.status == STATUS_FALHOU and a.erro_motivo == "leitura: RespostaInvalida"
    assert a.tokens_in == 18000 and a.custo_usd > 0


def test_reler_leitura_antiga_gera_o_formato_novo(db_session, empresa, teto):
    aid = _criar(db_session, empresa)
    a = db_session.get(AvaliadorAnalise, aid)  # leitura gravada no formato 1
    a.prompt_versao, a.status, a.resumo = "avaliador-v14-pilar", STATUS_PRONTA, "Resumo antigo."
    a.itens_json = json.dumps([{"subpilar": "D1", "tipo": "forte", "texto": "x", "evidencias": []}])
    db_session.commit()
    assert servico.formato_da_leitura(a) == servico.FORMATO_ITENS
    servico.refazer(db_session, a)
    db_session.commit()
    a = _rodar(db_session, aid, Dublê())
    assert a.prompt_versao == "avaliador-v27-marcas"
    assert servico.formato_da_leitura(a) == servico.FORMATO_MARCAS and a.resumo is None


def test_gate_cobra_as_cinco_chamadas(db_session, empresa, monkeypatch):
    from src.avaliador import marcas

    cobrado = {}
    monkeypatch.setattr(servico, "checar_teto", lambda s, e, est: cobrado.setdefault("usd", est))
    config_avaliador(monkeypatch, AVALIADOR_TETO_MENSAL_USD=10.0)
    a = db_session.get(AvaliadorAnalise, _criar(db_session, empresa))
    from src.avaliador.extracao import Documento

    regua = servico.carregar_regua()
    sistema, usuario = servico._pedido(
        regua, a.origem, a.descricao, Documento.de_json(a.texto_json)
    )
    assert cobrado["usd"] == pytest.approx(marcas.estimar_custo_usd(sistema, usuario, a.modelo, 4))


def test_pagina_da_web_vira_analise_com_o_endereco_lido(db_session, empresa, teto, monkeypatch):
    from tests.test_avaliador_pagina_web import HTML, Resp, Sessao, _resolve_para

    _resolve_para(monkeypatch, "93.184.216.34")
    s = Sessao(
        {
            "https://exemplo.com.br": Resp(301, headers={"Location": "/inicio"}),
            "https://exemplo.com.br/inicio": Resp(200, HTML.encode()),
        }
    )
    a = servico.criar_analise(
        db_session,
        empresa.id,
        None,
        DESCRICAO,
        cliente_nomeado=False,
        url="exemplo.com.br",
        sessao_http=s,
    )
    assert a.origem == "pagina" and a.formato == "web" and a.n_paginas == 4  # com o rodapé
    assert a.arquivo_nome == "https://exemplo.com.br/inicio"  # o que foi LIDO
    from src.avaliador.extracao import Documento

    doc = Documento.de_json(a.texto_json)
    sistema, _ = servico._pedido(servico.carregar_regua(), a.origem, a.descricao, doc)
    assert "COMO CONFERIR UMA MARCA" in sistema and "Há canal, não há horário." not in sistema


def test_duas_entradas_ao_mesmo_tempo_sao_recusadas(db_session, empresa, teto):
    with pytest.raises(AvaliacaoRecusada, match="um só"):
        servico.criar_analise(
            db_session,
            empresa.id,
            None,
            DESCRICAO,
            cliente_nomeado=False,
            texto="Cliente: oi",
            url="exemplo.com",
        )


def test_leitura_v20_na_fila_e_lida_com_o_pedido_de_hoje(db_session, empresa, teto):
    aid = _criar(db_session, empresa)
    db_session.get(AvaliadorAnalise, aid).prompt_versao = "avaliador-v20-veredito"
    db_session.commit()
    cli = Dublê()
    a = _rodar(db_session, aid, cli)
    assert a.status == STATUS_PRONTA and a.prompt_versao == "avaliador-v27-marcas"
    assert len(_leitura(cli.chamadas)) == 5


def test_peca_de_medicao_nao_perde_texto_nem_leitura(db_session, empresa, teto, monkeypatch):
    """21/10: as leituras que sustentam o gabarito das marcas não se relêem, não se concluem
    e não se apagam, e o watchdog não apaga o texto delas. A vizinha comum, sim."""
    agora = datetime.utcnow()
    fixa = _rodar(db_session, _criar(db_session, empresa), Dublê())
    comum = _rodar(db_session, _criar(db_session, empresa), Dublê())
    monkeypatch.setattr(servico, "PECAS_DE_MEDICAO", frozenset({fixa.id}))
    itens = fixa.itens_json
    for acao in (
        lambda: servico.refazer(db_session, fixa),
        lambda: servico.concluir(fixa),
        lambda: servico.apagar_leitura(fixa),
    ):
        with pytest.raises(AvaliacaoRecusada, match="Peça de medição"):
            acao()
    assert fixa.texto_json is not None and fixa.itens_json == itens
    fixa.usado_em = comum.usado_em = agora - timedelta(days=30)
    db_session.commit()
    assert servico.housekeeping(db_session, agora)["textos_apagados"] == 1
    assert fixa.texto_json is not None and comum.texto_json is None
    assert servico.PECAS_DE_MEDICAO != frozenset()  # o patch é deste teste


def test_pecas_de_medicao_de_producao():
    assert servico.PECAS_DE_MEDICAO == frozenset({24, 25, 26, 28, 31})
