"""O parecer no app (fatia 2, 22/10): campos de texto na tela da leitura, montagem a partir da
tela já corrigida, pendências que bloqueiam o Word, e o Word. Cliente LLM dublado, sem gasto."""

from __future__ import annotations

import io
import json

import pytest

from src.avaliador import parecer as P
from src.avaliador import servico
from src.avaliador.servico import AvaliacaoRecusada
from src.models.avaliador import AvaliadorAnalise, AvaliadorParecer
from src.models.empresa import Empresa
from tests.test_avaliador_servico import (
    HORARIO,
    REFERENCIA,
    Dublê,
    _criar,
    _rodar,
    config_avaliador,
)
from tests.test_ui_avaliador import _pronta, _url

SUBS = servico.SUBPILARES_ORDEM


@pytest.fixture
def empresa(db_session):
    e = Empresa(nome="Empresa Aval parecer")
    db_session.add(e)
    db_session.commit()
    return e


@pytest.fixture
def teto(monkeypatch):
    config_avaliador(monkeypatch, AVALIADOR_TETO_MENSAL_USD=10.0)


def _pronta_servico(db_session, empresa):
    return _rodar(db_session, _criar(db_session, empresa), Dublê())


def _montar(a):
    tela = servico.organizar_marcas(servico.carregar_regua(), a, servico.correcoes_de(a))
    return P.montar(tela, servico.textos_do_parecer(a), a.id, "22/10/2026",
                    json.loads(a.itens_json)["cobertura"])  # fmt: skip


def _preencher(db_session, a):
    servico.salvar_parecer(db_session, a, "empresa_da_peca", "Empresa da Peça SA", None)
    servico.salvar_parecer(db_session, a, "ordem", "A ruptura está na Precisão.", None)
    for c in SUBS:
        servico.salvar_parecer(db_session, a, c, f"Achado de {c}.", None)
    db_session.commit()


def test_montar_escolhe_um_trecho_por_marca_e_lista_o_que_falta(db_session, empresa, teto):
    a = _pronta_servico(db_session, empresa)
    p = _montar(a)
    d1 = next(s for pl in p["pilares"] for s in pl["subpilares"] if s["codigo"] == "D1")
    # D1.01 e D1.04 citam o mesmo trecho: entra uma vez só. O trecho começa no meio da linha
    # da peça, e o parecer imprime a linha inteira (23/10)
    assert d1["diz"] == ["Atendimento de segunda a sexta das 8h às 18h pelo gerente de conta."]
    assert "Não diz em que dias da semana há expediente." in d1["nao_diz"]
    assert not any("D1." in t for t in d1["nao_diz"])  # código nunca aparece
    a1 = next(s for pl in p["pilares"] for s in pl["subpilares"] if s["codigo"] == "A1")
    assert a1["diz"] == [REFERENCIA[1]]
    p1 = next(s for pl in p["pilares"] for s in pl["subpilares"] if s["codigo"] == "P1")
    assert p1["diz"] == []  # vai para o Word como "Não encontramos no texto…"
    assert [pl["nome"] for pl in p["pilares"]] == [
        "Precisão",
        "Disponibilidade",
        "Parceria",
        "Aconselhamento",
    ]
    assert p["pendencias"][:2] == [
        "Escreva o nome da empresa da peça, que sai no título do Word.",
        "Escreva a frase da ordem, no alto.",
    ]
    assert sum("escreva o achado" in x for x in p["pendencias"]) == 12
    assert p["cabecalho"].startswith(f"Leitura nº {a.id}, 22/10/2026. Lidas")


def test_ate_quatro_trechos_na_ordem_do_manual():
    marcas = [
        {
            "codigo": f"P1.0{i}",
            "estado": "esta",
            "evidencias": [{"trecho": f"trecho número {i} da peça", "conferido": True}],
        }
        for i in range(1, 7)
    ]
    assert [e["codigo"] for e in P._escolhidos(marcas)] == ["P1.01", "P1.02", "P1.03", "P1.04"]


def test_salvar_campos_e_pendencias_somem(db_session, empresa, teto):
    a = _pronta_servico(db_session, empresa)
    _preencher(db_session, a)
    assert db_session.query(AvaliadorParecer).count() == 1
    assert servico.textos_do_parecer(a)["achados"]["Pa2"] == "Achado de Pa2."
    assert _montar(a)["pendencias"] == []
    with pytest.raises(AvaliacaoRecusada, match="desconhecido"):
        servico.salvar_parecer(db_session, a, "outra coisa", "x", None)


def test_sem_conferencia_bloqueia_ate_ser_resolvida(db_session, empresa, teto):
    a = _pronta_servico(db_session, empresa)
    _preencher(db_session, a)
    g = json.loads(a.itens_json)
    m = next(m for m in g["marcas"] if m["codigo"] == "D1.05")
    m["esta"], m["evidencias"] = True, [{"pagina": 2, "trecho": "inventado", "conferido": False}]
    a.itens_json = json.dumps(g)
    db_session.commit()
    assert any("resolva a marca sem conferência" in x for x in _montar(a)["pendencias"])
    servico.recusar_nao_achadas(db_session, a, "D1.05", None)
    db_session.commit()
    p = _montar(a)
    assert p["pendencias"] == []
    d1 = next(s for pl in p["pilares"] for s in pl["subpilares"] if s["codigo"] == "D1")
    assert "Não diz em que dias da semana há expediente." in d1["nao_diz"]
    with pytest.raises(AvaliacaoRecusada, match="não tem citação sem conferência"):
        servico.recusar_nao_achadas(db_session, a, "D1.01", None)


def test_correcao_entra_no_parecer(db_session, empresa, teto):
    a = _pronta_servico(db_session, empresa)
    servico.recusar_trecho(db_session, a, "A1.02", REFERENCIA[1], None)
    db_session.commit()
    a1 = next(s for pl in _montar(a)["pilares"] for s in pl["subpilares"] if s["codigo"] == "A1")
    assert (
        a1["diz"] == []
        and "Não traz reconhecimento de terceiro: publicação, prêmio, parceiro." in a1["nao_diz"]
    )


def test_apagar_leitura_apaga_o_texto_do_parecer(db_session, empresa, teto):
    a = _pronta_servico(db_session, empresa)
    _preencher(db_session, a)
    servico.apagar_leitura(a)
    db_session.commit()
    p = db_session.query(AvaliadorParecer).one()
    assert (p.ordem, p.nota_c, p.achados_json) == ("", "", "{}")


def test_tela_e_word(client_loyall, db_session, empresa, teto):
    aid = _pronta(client_loyall, db_session, empresa)
    html = client_loyall.get(_url(empresa, f"/{aid}")).get_data(as_text=True)
    assert 'data-secao="parecer"' in html and "A ordem" in html and "O achado" in html
    assert (
        "vai para o parecer" in html
        and "Para o Word sair, falta: o nome da empresa da peça, a frase da ordem, 12 achados."
        in html
    )
    r = client_loyall.get(_url(empresa, f"/{aid}/parecer.docx"))
    assert r.status_code == 409 and "Escreva a frase da ordem" in r.get_data(as_text=True)
    salvo = client_loyall.post(
        _url(empresa, f"/{aid}/parecer"), data={"campo": "ordem", "texto": "A ordem."}
    ).get_data(as_text=True)
    assert (
        "salvo" in salvo
        and 'hx-swap-oob="true"' in salvo
        and "Escreva a frase da ordem" not in salvo
    )
    for c in SUBS:
        client_loyall.post(
            _url(empresa, f"/{aid}/parecer"), data={"campo": c, "texto": f"Achado {c}."}
        )
    client_loyall.post(
        _url(empresa, f"/{aid}/parecer"), data={"campo": "empresa_da_peca", "texto": "Peça SA"}
    )
    r = client_loyall.get(_url(empresa, f"/{aid}/parecer.docx"))
    assert r.status_code == 200 and r.mimetype.endswith("wordprocessingml.document")
    from docx import Document

    texto = "\n".join(p.text for p in Document(io.BytesIO(r.data)).paragraphs)
    assert (
        "Parecer PDPA: Peça SA" in texto
        and "Empresa Aval parecer" not in texto  # nunca o nome da empresa da conta
        and "A ordem." in texto
        and "Achado Pa2." in texto
    )
    assert "Não encontramos no texto nada que responda a essa expectativa." in texto
    assert "D1.01" not in texto and "—" not in texto
    assert (
        client_loyall.post(
            _url(empresa, f"/{aid}/parecer"), data={"campo": "x", "texto": "y"}
        ).status_code
        == 200
    )
    a = db_session.get(AvaliadorAnalise, aid)
    assert a.status == "pronta"


def _duas(db_session, a, codigo, t1, t2):
    g = json.loads(a.itens_json)
    m = next(m for m in g["marcas"] if m["codigo"] == codigo)
    m["esta"] = True
    m["evidencias"] = [
        {"pagina": 2, "trecho": t, "conferido": True, "pagina_encontrada": 2} for t in (t1, t2)
    ]
    a.itens_json = json.dumps(g)
    db_session.commit()


def _encher_p1(db_session, a):
    """P1.02, P1.03 e P1.04 respondidas, um trecho cada: com a P1.01 o bloco fecha os quatro
    lugares na primeira volta, e o segundo trecho da P1.01 fica de fora (sobra o botão)."""
    g = json.loads(a.itens_json)
    frases = {
        "P1.02": "Não estão incluídos deslocamentos fora da capital.",
        "P1.03": "A instalação acontece em até dez dias úteis após o aceite.",
        "P1.04": "O cliente envia a planta baixa antes da visita técnica.",
    }
    for m in g["marcas"]:
        if m["codigo"] in frases:
            m["esta"] = True
            m["evidencias"] = [
                {
                    "pagina": 2,
                    "trecho": frases[m["codigo"]],
                    "conferido": True,
                    "pagina_encontrada": 2,
                }
            ]
    a.itens_json = json.dumps(g)
    db_session.commit()


def test_usar_este_no_parecer_troca_o_trecho_sem_virar_correcao(db_session, empresa, teto):
    from src.models.avaliador import AvaliadorCorrecao

    a = _pronta_servico(db_session, empresa)
    _duas(
        db_session,
        a,
        "P1.01",
        "Trecho curto da marca",
        "Outro trecho, bem mais comprido, da mesma marca",
    )
    _encher_p1(db_session, a)
    p1 = lambda: next(  # noqa: E731
        s for pl in _montar(a)["pilares"] for s in pl["subpilares"] if s["codigo"] == "P1"
    )
    assert p1()["diz"][0] == "Trecho curto da marca" and len(p1()["diz"]) == 4
    servico.escolher_trecho(
        db_session, a, "P1.01", "Outro trecho, bem mais comprido, da mesma marca", None
    )
    db_session.commit()
    assert p1()["diz"][0] == "Outro trecho, bem mais comprido, da mesma marca"
    assert "Trecho curto da marca" not in p1()["diz"]
    assert db_session.query(AvaliadorCorrecao).count() == 0  # a tabela de erros fica limpa
    # o escolhido foi recusado depois: volta ao automático
    servico.recusar_trecho(
        db_session, a, "P1.01", "Outro trecho, bem mais comprido, da mesma marca", None
    )
    db_session.commit()
    assert p1()["diz"][0] == "Trecho curto da marca"
    with pytest.raises(AvaliacaoRecusada, match="não é uma resposta desta marca"):
        servico.escolher_trecho(
            db_session, a, "P1.01", "Outro trecho, bem mais comprido, da mesma marca", None
        )


def test_tela_mostra_usar_este_so_em_marca_que_vai_para_o_parecer(
    client_loyall, db_session, empresa, teto
):
    aid = _pronta(client_loyall, db_session, empresa)
    a = db_session.get(AvaliadorAnalise, aid)
    _duas(
        db_session,
        a,
        "P1.01",
        "Trecho curto da marca",
        "Outro trecho, bem mais comprido, da mesma marca",
    )
    _encher_p1(db_session, a)
    html = client_loyall.get(_url(empresa, f"/{aid}")).get_data(as_text=True)
    assert html.count("usar este no parecer") == 1  # só o segundo da P1.01
    html = client_loyall.post(
        _url(empresa, f"/{aid}/corrigir/escolher"),
        data={"codigo": "P1.01", "trecho": "Outro trecho, bem mais comprido, da mesma marca"},
    ).get_data(as_text=True)
    marca = html.split('data-marca="P1.01"', 1)[1].split("</details>", 1)[0]
    # fechada, a marca mostra o trecho que vai para o parecer
    fechada, bloco = marca.split("</summary>", 1)
    assert "no parecer:" in fechada and "Outro trecho" in fechada and "Trecho curto" not in fechada
    # aberta, depois da troca: o primeiro trecho ganha o botão, o segundo vai para o parecer
    marcos = ("Trecho curto", "usar este no parecer", "Outro trecho", "vai para o parecer")
    ordem = [bloco.index(x) for x in marcos]
    assert ordem == sorted(ordem)


def _tela(blocos):
    """Tela mínima: [(subpilar, [(código, [(trecho, referência[, seção])])])], tudo 'está'."""
    from types import SimpleNamespace as N

    return {
        "pilares": [
            {
                "subpilares": [
                    {
                        "subpilar": N(codigo=sp),
                        "marcas": [
                            {
                                "codigo": c,
                                "estado": "esta",
                                "evidencias": [
                                    {
                                        "trecho": t[0],
                                        "conferido": True,
                                        "referencia": t[1],
                                        "pagina_encontrada": t[2] if len(t) > 2 else 1,
                                    }
                                    for t in ts
                                ],
                            }
                            for c, ts in marcas
                        ],
                    }
                    for sp, marcas in blocos
                ]
            }
        ]
    }


KM = "Importante: não há opção de quilometragem ilimitada."


def test_trecho_impresso_nao_volta_a_nao_ser_que_o_bloco_fique_vazio():
    tela = _tela(
        [
            ("P1", [("P1.02", [(KM, "")])]),
            (
                "D3",
                [("D3.01", [("O atraso pode impactar sua reserva.", "")]), ("D3.04", [(KM, "")])],
            ),
            ("A2", [("A2.07", [(KM, "")])]),  # só tem o KM: o bloco ficaria vazio
        ]
    )
    sel = P.selecao(tela)
    assert [e["trecho"] for e in sel["P1"]] == [KM]
    assert [e["trecho"] for e in sel["D3"]] == ["O atraso pode impactar sua reserva."]
    assert [e["trecho"] for e in sel["A2"]] == [KM]
    # o escolhido pela pessoa manda, mesmo já impresso antes
    sel = P.selecao(tela, {"D3.04": KM})
    assert KM in [e["trecho"] for e in sel["D3"]]


def test_rodape_por_ultimo_e_o_resto_na_ordem_da_peca():
    rod = "página /x, rodapé da página"
    tela = _tela(
        [
            (
                "D1",
                [
                    (
                        "D1.04",
                        [
                            ("Central 24h: 0800 979", rod, 9),
                            ("Atendimento das 8h às 18h, de segunda a sexta", "seção Contato", 4),
                        ],
                    )
                ],
            ),
            (
                "A1",
                [
                    (
                        "A1.04",
                        [
                            ("Veículo: AGRO+", "página /media/", 6),
                            (
                                "A E2Carbon e o Sicredi uniram forças no campo.",
                                "página /sicredi/",
                                5,
                            ),
                        ],
                    )
                ],
            ),
            ("A2", [("A2.01", [("Primeiro na leitura, mais longo", "", 3), ("Curto", "", 3)])]),
        ]
    )
    sel = P.selecao(tela)
    assert sel["D1"][0]["trecho"].startswith("Atendimento")  # o do rodapé fica por último
    # o rótulo é o menor, mas vem depois na peça (a #32, 22/10)
    assert sel["A1"][0]["trecho"].startswith("A E2Carbon e o Sicredi")
    assert sel["A2"][0]["trecho"] == "Primeiro na leitura, mais longo"  # mesma seção: leitura


def _m(codigo, estado, nome, exc=False):
    return {"codigo": codigo, "estado": estado, "nome_curto": nome, "texto": "Diz x?",
            "so_cliente_nomeado": exc}  # fmt: skip


def test_linha_da_frente_so_com_nome_curto_sem_contagem():
    """Página da frente (22/10): o que a peça diz e o que falta, pelo nome curto do Manual."""
    ms = [
        _m("P2.01", "nao_esta", "quando termina"),
        _m("P2.02", "nao_esta", "padrão"),
        _m("P2.03", "nao_esta", "critério de aceitação"),
    ]
    assert P.linha_da_frente(ms, False) == (
        "Nada escrito. Falta: quando termina, padrão e critério de aceitação."
    )
    ms = [
        _m("P1.01", "esta", "o que entrega"),
        _m("P1.02", "esta", "o que não entrega"),
        _m("P1.07", "nao_esta", "quem assina"),
        _m("P1.10", "sem_conferencia", "divisão de responsabilidade"),
    ]
    assert P.linha_da_frente(ms, False) == (
        "Diz: o que entrega e o que não entrega. Falta: quem assina. "
        "A conferir: divisão de responsabilidade."
    )
    assert P.linha_da_frente([_m("D1.01", "esta", "por onde falar")], None) == (
        "Diz: por onde falar. Nada falta."
    )


def test_linha_da_frente_excecao_so_com_cliente_nomeado():
    ms = [_m("Pa2.01", "nao_esta", "preço"), _m("Pa2.03", "nao_esta", "impostos", exc=True)]
    assert P.linha_da_frente(ms, True) == "Nada escrito. Falta: preço e impostos."
    for sem in (False, None):  # sem cliente, ou sem a resposta: a exceção fica fora
        assert P.linha_da_frente(ms, sem) == "Nada escrito. Falta: preço."


def test_frente_na_tela_e_no_word(client_loyall, db_session, empresa, teto):
    aid = _pronta(client_loyall, db_session, empresa)
    a = db_session.get(AvaliadorAnalise, aid)
    p = _montar(a)
    assert [x["codigo"] for x in p["frente"]] == list(SUBS)
    d1 = next(x for x in p["frente"] if x["codigo"] == "D1")
    assert d1["pergunta"] == "Quando eu precisar falar, eu consigo?"
    assert d1["linha"].startswith("Diz: ") and "dias de expediente" in d1["linha"]
    html = client_loyall.get(_url(empresa, f"/{aid}")).get_data(as_text=True)
    assert "Em uma página (a frente do Word)" in html and d1["linha"] in html
    _preencher(db_session, a)
    r = client_loyall.get(_url(empresa, f"/{aid}/parecer.docx"))
    from docx import Document

    texto = "\n".join(x.text for x in Document(io.BytesIO(r.data)).paragraphs)
    assert texto.index("Em uma página") < texto.index("A ordem")
    # pergunta e linha no mesmo parágrafo, separadas por quebra (cabe numa página)
    pars = Document(io.BytesIO(r.data)).paragraphs
    par = next(x for x in pars if x.text.startswith("Acessibilidade · "))
    assert par.text == "Acessibilidade · Quando eu precisar falar, eu consigo?\n" + d1["linha"]
    assert par._p.xml.count("<w:br/>") == 1


def test_titulo_do_word_e_a_empresa_da_peca_nunca_a_da_conta(
    client_loyall, db_session, empresa, teto
):
    """22/10: o Word da #32 (E2Carbon) saiu com o nome da empresa da conta em que a leitura
    foi feita. O título é o nome escrito no parecer; em branco, o Word não sai."""
    aid = _pronta(client_loyall, db_session, empresa)
    a = db_session.get(AvaliadorAnalise, aid)
    _preencher(db_session, a)
    servico.salvar_parecer(db_session, a, "empresa_da_peca", "", None)
    db_session.commit()
    r = client_loyall.get(_url(empresa, f"/{aid}/parecer.docx"))
    assert r.status_code == 409 and "nome da empresa da peça" in r.get_data(as_text=True)
    servico.salvar_parecer(db_session, a, "empresa_da_peca", "  Prospect Lida SA ", None)
    db_session.commit()
    r = client_loyall.get(_url(empresa, f"/{aid}/parecer.docx"))
    from docx import Document

    d = Document(io.BytesIO(r.data))
    texto = "\n".join(x.text for x in d.paragraphs) + d.core_properties.title
    assert d.paragraphs[0].text == "Parecer PDPA: Prospect Lida SA"
    assert empresa.nome not in texto  # o nome da conta não vai ao cliente
    # apagar a leitura tira o nome junto com o resto do parecer
    servico.apagar_leitura(a)
    db_session.commit()
    assert db_session.query(AvaliadorParecer).one().empresa_da_peca == ""


def test_trecho_com_a_maior_parte_em_comum_nao_se_repete():
    """22/10, Exemplo da #32: dois trechos que se sobrepõem sem um conter o outro."""
    t1 = (
        "Na primeira etapa do programa, dez produtores certificados pelo Fair Trade, ligados à "
        "Cooperativa dos Costas, em Boa Esperança (MG), participaram primeira etapa do projeto."
    )
    t2 = (
        "dez produtores certificados pelo Fair Trade, ligados à Cooperativa dos Costas, em Boa "
        "Esperança (MG), participaram primeira etapa do projeto. São pequenas propriedades de "
        "agricultura familiar"
    )
    outro = "A E2Carbon e o Sicredi uniram forças para impulsionar a sustentabilidade no campo."
    tela = _tela([("A1", [("A1.04", [(t1, "", 5)]), ("A1.06", [(t2, "", 5), (outro, "", 5)])])])
    assert [e["trecho"] for e in P.selecao(tela)["A1"]] == [t1, outro]
    # pouco texto em comum não é repetição, nem começo igual curto
    assert not P._repete("o carro reserva chega em 24h", "o carro novo chega amanhã cedo")
    assert not P._repete("veiculo: agro+", "veiculo: embrapii")


def test_completa_ate_quatro_com_o_segundo_trecho_das_marcas():
    """23/10, Fibrav: duas marcas respondidas com três trechos cada saíam com dois trechos no
    Word. Primeira volta, um de cada marca; depois o seguinte, até quatro."""
    tela = _tela(
        [
            (
                "P2",
                [
                    ("P2.02", [("Norma ASME RTP-1 de vasos.", "", 3),
                               ("Ensaio de dureza Barcol em cada lote.", "", 3),
                               ("Perda por ignição medida em laboratório.", "", 3)]),
                    ("P2.03", [("Ensaios de tração e flexão do laminado.", "", 4),
                               ("Espessura mínima conferida na entrega.", "", 4),
                               ("Teste hidrostático antes do embarque.", "", 4)]),
                ],
            )
        ]
    )  # fmt: skip
    assert [(e["codigo"], e["trecho"][:12]) for e in P.selecao(tela)["P2"]] == [
        ("P2.02", "Norma ASME R"),
        ("P2.03", "Ensaios de t"),
        ("P2.02", "Ensaio de du"),
        ("P2.03", "Espessura mí"),
    ]


# ── a linha inteira do trecho que começa no meio dela (23/10) ──────────────────

LINHA_31 = (
    "Você pode devolver o carro em qualquer agência Localiza. "
    "Se for diferente da que você abriu, haverá taxa de retorno."
)


def _doc(formato, texto):
    from src.avaliador.extracao import Documento, Pagina

    return Documento(formato, [Pagina(1, texto)])


@pytest.mark.parametrize(
    "formato, texto, trecho, sai",
    [
        # começa no meio: a linha inteira
        ("web", f"Devolução\n{LINHA_31}", "Se for diferente da que você abriu", LINHA_31),
        # começa no começo da linha, ou depois só de marcador: o trecho como foi citado
        ("web", f"Devolução\n{LINHA_31}", "Você pode devolver o carro em qualquer", None),
        ("web", "• Atendimento por telefone e e-mail", "Atendimento por telefone", None),
        # título emendado ao texto de baixo (a emenda aceita na conferência): começa na linha
        ("web", f"Devolução\n{LINHA_31}", "Devolução: Você pode devolver o carro", None),
        # PDF: a quebra simples é a linha impressa; a unidade é o parágrafo
        (
            "pdf",
            "Você pode devolver o carro em qualquer\nagência Localiza. Se for diferente da que\n"
            "você abriu, haverá taxa de retorno.\n\nOutro parágrafo.",
            "Se for diferente da que você abriu",
            LINHA_31,
        ),
        # tabela: a célula é a unidade
        ("pptx", "Plano | Inclui suporte técnico 24h | R$ 10", "Inclui suporte técnico", None),
        (
            "pptx",
            "Plano | Inclui suporte técnico 24h | R$ 10",
            "suporte técnico 24h",
            "Inclui suporte técnico 24h",
        ),  # fmt: skip
        # não achado (texto diferente): o trecho como foi citado
        ("web", LINHA_31, "Isto não está na peça de jeito nenhum", None),
    ],
)
def test_linha_inteira_so_quando_o_trecho_comeca_no_meio(formato, texto, trecho, sai):
    from src.avaliador.extracao import linha_inteira

    assert linha_inteira(_doc(formato, texto), 1, trecho) == sai


def test_linha_longa_sai_com_o_trecho_e_reticencias():
    from src.avaliador.extracao import LIMITE_LINHA, linha_inteira

    longa = "Texto de antes. " * 40 + "Se for diferente da que você abriu, há taxa."
    assert len(longa) > LIMITE_LINHA
    assert (
        linha_inteira(_doc("web", longa), 1, "Se for diferente da que você abriu")
        == "… Se for diferente da que você abriu"
    )


def test_sem_o_texto_da_peca_sai_o_trecho_como_foi_citado():
    e = {"trecho": "Se for diferente da que você abriu", "pagina_encontrada": 1}
    assert P.impresso(e, None) == e["trecho"]
    assert P.impresso(e, _doc("web", LINHA_31)) == LINHA_31


def test_dois_trechos_da_mesma_linha_saem_uma_vez_so():
    tela = _tela(
        [
            ("P1", [("P1.01", [("Você pode devolver o carro em qualquer agência", "")]),
                    ("P1.02", [("Se for diferente da que você abriu", "")])]),
        ]
    )  # fmt: skip
    tela["documento"] = _doc("web", f"Devolução\n{LINHA_31}")
    sel = P.selecao(tela)
    assert [P.impresso(e, tela["documento"]) for e in sel["P1"]] == [
        "Você pode devolver o carro em qualquer agência"
    ]
    tela["documento"] = None  # sem a linha inteira os dois trechos são diferentes
    for x in tela["pilares"][0]["subpilares"][0]["marcas"]:
        for e in x["evidencias"]:
            e.pop("impresso", None)
    assert len(P.selecao(tela)["P1"]) == 2


def test_word_e_rascunho_levam_a_linha_inteira_a_tela_e_a_base_nao(
    client_loyall, db_session, empresa, teto
):
    from src.avaliador import rascunho as R

    a = _pronta_servico(db_session, empresa)
    tela = servico.organizar_marcas(servico.carregar_regua(), a, servico.correcoes_de(a))
    linha = "Atendimento de segunda a sexta das 8h às 18h pelo gerente de conta."
    texto, base = R.entrada(tela, {}, a.descricao)
    assert linha in texto
    # a base é a dos trechos como foram citados: igual à de uma tela sem o texto da peça
    sem_doc = dict(tela, documento=None)
    for b in sem_doc["pilares"]:
        for sp in b["subpilares"]:
            for x in sp["marcas"]:
                for e in x["evidencias"]:
                    e.pop("impresso", None)
    texto_sem, base_sem = R.entrada(sem_doc, {}, a.descricao)
    assert linha not in texto_sem and HORARIO[1] in texto_sem and base == base_sem
    _preencher(db_session, a)
    p = _montar(a)
    assert not p["pendencias"]
    from docx import Document

    pars = [x.text for x in Document(io.BytesIO(P.gerar_docx(p))).paragraphs]
    assert f"“{linha}”" in pars
    html = client_loyall.get(_url(empresa, f"/{a.id}")).get_data(as_text=True)
    assert HORARIO[1] in html and linha not in html  # a tela mostra o trecho citado
