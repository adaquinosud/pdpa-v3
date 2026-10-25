"""Marcas de leitura de documento (capítulo 17 do Manual) como dado da régua.

Fatia 1 do balizador (18/10): carregadas e validadas, ainda sem uso na leitura. Nenhuma
chamada paga."""

from __future__ import annotations

import copy
import importlib.util
from pathlib import Path

import pytest

from src.api.painel import SUBPILARES_ORDEM
from src.avaliador.regua import (
    GRUPO_FORA,
    MARCAS_APOSENTADAS,
    Regua,
    ReguaInvalida,
    _ler,
    carregar_regua,
    validar,
)

DOCX = Path("data/PDPA_Manual_Operacao_v16.docx")


def _script():
    spec = importlib.util.spec_from_file_location("x", "scripts/extrair_marcas.py")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def _regua(marcas):
    return Regua(
        _ler("manual_cap2.json"), _ler("frases_l1.json"), _ler("frases_l3.json"), "t", marcas
    )


@pytest.mark.skipif(not DOCX.exists(), reason="Manual .docx não está nesta máquina")
def test_extrato_bate_com_o_docx_do_manual():
    """Fonte única: o Manual. Mudou o capítulo → rode scripts/extrair_marcas.py."""
    assert _script().extrair(DOCX) == _ler("marcas.json")


def test_marcas_como_decidido_em_18_10_e_na_regua_de_79():
    """18/10: 87 marcas, 7 exceções. 21/10 (Manual v13): oito saem, sem renumerar, e três
    exceções passam a valer sempre (P1.07, Pa2.04, Pa2.09)."""
    r = carregar_regua()
    por_grupo = {g: len(r.marcas_de(g)) for g in SUBPILARES_ORDEM + [GRUPO_FORA]}
    assert por_grupo == {
        "P1": 7, "P2": 7, "P3": 3, "D1": 8, "D2": 5, "D3": 5,
        "Pa1": 8, "Pa2": 8, "Pa3": 5, "A1": 6, "A2": 8, "A3": 5, "C": 4,
    }  # fmt: skip
    assert len(r.marcas) == 79
    assert [m.codigo for m in r.marcas_de("P1")] == [
        "P1.01", "P1.02", "P1.03", "P1.04", "P1.07", "P1.10", "P1.11",
    ]  # fmt: skip
    assert not {m.codigo for m in r.marcas} & MARCAS_APOSENTADAS
    assert [m.codigo for m in r.marcas if m.so_cliente_nomeado] == [
        "Pa1.05", "Pa1.06", "Pa1.07", "Pa2.03",
    ]  # fmt: skip
    assert all(m.porque_importa == "" for m in r.marcas)  # nasce vazio, de propósito
    assert r.marcas_de("A3")[3].texto == "Recomenda comprar menos, ou depois, ou outra coisa?"
    assert r.versao.startswith("manual+marcas+pistas:")


def test_textos_do_manual_de_hoje():
    """Os sete textos trocados em 21/10 (régua de 79), os de 21/10 no v12 que ficaram, e o lote
    de 24/10 no v15: sobem P2.04, Pa1.06, Pa3.01 e Pa3.02; P2.01, Pa2.02, A3.04 e C.04 não
    passaram no portão e ficam com o texto de antes (a Pa2.02 voltou pela exceção de 24/10, no v16).
    """
    texto = {m.codigo: m.texto for m in carregar_regua().marcas}
    assert texto["P2.01"] == "Diz em que momento a entrega está concluída?"
    # P2.02: "ou método" saiu no portão (balançou numa de cinco leituras); fica o texto do v12
    assert texto["P2.02"] == "Diz qual padrão, protocolo ou referência ela segue?"
    assert texto["P2.04"] == (
        "Diz o que a empresa devolve em dinheiro, refaz ou desconta se não cumprir o que prometeu?"
    )
    assert texto["P3.04"] == (
        "Diz o que acontece se o cliente precisar de mais ou de menos do que foi combinado?"
    )
    assert texto["Pa1.06"] == "Diz quais problemas a empresa que vai ler esta peça está vivendo?"
    assert texto["Pa3.01"] == "Diz o que a empresa faz pelo cliente depois que o contrato acaba?"
    assert texto["Pa3.02"] == (
        "Diz por quanto tempo depois do fim do contrato a empresa continua disponível?"
    )
    # Pa2.02: a exceção de 24/10 (§4.73), no v16: sem a lista de alternativas
    assert texto["Pa2.02"] == "Diz de que depende o valor que será cobrado?"
    assert texto["A3.04"] == "Recomenda comprar menos, ou depois, ou outra coisa?"
    assert texto["C.04"] == "Nomeia as tecnologias, softwares ou ferramentas que usa?"
    curto = {m.codigo: m.nome_curto for m in carregar_regua().marcas}
    assert curto["Pa1.06"] == "problemas que o cliente vive"
    assert curto["Pa3.01"] == "depois do contrato"
    assert (
        curto["A3.04"] == "recomendação de comprar menos" and curto["C.04"] == "tecnologia que usa"
    )
    assert texto["Pa1.08"] == "Oferece uma forma de experimentar antes de pagar?"
    assert texto["A1.01"] == "Traz alguém de fora falando na peça, com nome e cargo?"
    assert texto["A1.03"] == (
        "Traz prova de fora de que sabe fazer: certificação, registro, selo ou formação de"
        " quem executa?"
    )
    assert texto["P1.11"].startswith("Diz o que acontece se o cliente pedir, depois de fechado")
    assert texto["Pa2.01"] == "Diz quanto custa, com valor em número?"
    assert texto["C.01"].startswith("Cita dado ou pesquisa nomeando")
    assert texto["C.02"].startswith("Cita dado ou pesquisa com a data")
    assert texto["Pa3.03"] == "Diz qual é o canal de atendimento depois da contratação?"


@pytest.mark.parametrize(
    "estraga, motivo",
    [
        (lambda d: d["marcas"][1].update(codigo="P1.01"), "marca repetida: P1.01"),
        (
            lambda d: d["marcas"].insert(0, d["marcas"].pop(3)),
            "numeração das marcas de P1 fora de ordem",
        ),
        (lambda d: d["marcas"][3].update(codigo="P1.05"), "marca P1.05 usa código aposentado"),
        (lambda d: d["marcas"][0].update(grupo="P2"), "marca P1.01 no grupo P2"),
        (lambda d: d["marcas"][0].update(grupo="X9"), "grupo fora do modelo: X9"),
        (lambda d: d["marcas"][0].update(texto=" "), "marca P1.01 sem texto"),
        (lambda d: d["marcas"][0].update(nome_curto=""), "marca P1.01 sem nome curto"),
        (lambda d: d["marcas"][1].update(nome_curto="o que entrega"), "nome curto repetido"),
        (lambda d: d["perguntas_leitor"].pop("D2"), r"sem pergunta do leitor: \['D2'\]"),
        (
            lambda d: d.update(marcas=[x for x in d["marcas"] if x["grupo"] != "D3"]),
            "nenhuma marca em D3",
        ),
        (lambda d: d["pilares"][0].update(titulo="Pilar 1 — PROMESSA"), "pilares das marcas"),
    ],
)
def test_marcas_fora_da_regra_derrubam_a_regua(estraga, motivo):
    d = copy.deepcopy(_ler("marcas.json"))
    estraga(d)
    with pytest.raises(ReguaInvalida, match=motivo):
        validar(_regua(d))


def test_extrator_recusa_capitulo_fora_da_forma(tmp_path):
    import docx

    mod = _script()

    def manual(*linhas):
        doc = docx.Document()
        for t in ("16. Outro capítulo", "texto", "17. Marcas de leitura de documento", *linhas):
            doc.add_paragraph(t)
        doc.add_paragraph("Sobre a Loyall")
        caminho = tmp_path / "m.docx"
        doc.save(caminho)
        return caminho

    base = (
        "Abertura.",
        "Pilar 1 — PRECISÃO",
        "Apresenta a promessa?",
        "P1 — Calibração da Promessa",
        "Pergunta do leitor: O que me venderam é o que eu vou receber?",
        "P1.01 Diz o que será entregue? (nome curto: o que entrega)",
    )
    exc = "Só se esperam em peça feita para um cliente com nome: P1.01."
    ok = mod.extrair(manual(*base, exc))
    assert ok["fonte"].endswith("capítulo 17") and ok["marcas"][0]["so_cliente_nomeado"]
    assert ok["abertura"] == ["Abertura."]
    # Manual v14: o nome curto sai do texto da marca (o modelo lê a marca sem ele)
    assert ok["marcas"][0]["texto"] == "Diz o que será entregue?"
    assert ok["marcas"][0]["nome_curto"] == "o que entrega"
    assert ok["perguntas_leitor"] == {"P1": "O que me venderam é o que eu vou receber?"}
    casos = [
        ((*base, "P2.01 Diz outra coisa? (nome curto: x)", exc), "fora do seu subpilar"),
        ((*base, "P1.02 Diz outra coisa?", exc), "P1.02 sem nome curto"),
        ((*base[:4], base[5], exc), r"subpilar sem pergunta do leitor: \['P1'\]"),
        ((*base, "Só se esperam em peça feita para um cliente com nome: P1.09."), "P1.09"),
        (base, "não tem o parágrafo"),
        ((*base, exc, "mais texto"), "depois das exceções"),
        (
            (*base[:3], "P1 — Promessa", "P1.01 Diz? (nome curto: x)", exc),
            "o capítulo diz 'Promessa'",
        ),
    ]
    for linhas, motivo in casos:
        with pytest.raises(SystemExit, match=motivo):
            mod.extrair(manual(*linhas))


def test_nome_curto_e_pergunta_do_leitor_nao_vao_para_o_modelo():
    """Manual v14 (22/10): o nome curto e a pergunta do leitor são da página da frente do
    parecer. O pedido da leitura continua só com o texto de cada marca."""
    from src.avaliador import marcas

    r = carregar_regua()
    # o pedido inteiro: o sistema e a instrução de cada chamada (as marcas vão nas instruções)
    sistema = marcas.montar_sistema(r) + "".join(
        marcas.instrucao_etapa(r, t, g) for t, g in marcas.etapas(r)
    )
    assert all(f"{m.codigo} {m.texto}\n" in sistema + "\n" for m in r.marcas)
    assert "nome curto" not in sistema
    for m in r.marcas:
        assert f"{m.texto} ({m.nome_curto})" not in sistema
    assert not any(q in sistema for q in r.perguntas_leitor.values())
